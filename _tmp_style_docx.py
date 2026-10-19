from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

path = r"C:\Users\7onch\Desktop\odluka - napojnice.docx"

doc = Document()

# Page margins close to the original document
section = doc.sections[0]
section.top_margin = Cm(2.5)
section.bottom_margin = Cm(2.5)
section.left_margin = Cm(2.5)
section.right_margin = Cm(2.5)

style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

lines = [
    "ZIA ROSA J.D.O.O. (zastupano po direktoru Suzana Blažeić)",
    "OIB: 69327780069",
    "Kumičićeva 55A",
    "51000 Rijeka",
    "",
    "Donosi",
]

for t in lines:
    p = doc.add_paragraph(t)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_after = Pt(6)

p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(12)

p = doc.add_paragraph("ODLUKU")
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = p.runs[0]
run.bold = True
p.paragraph_format.space_after = Pt(12)

tekst = (
    "da od 24.12.2024. djelatnici Zia Rosa j.d.o.o. i studenti koji rade putem "
    "studentskog servisa NE primaju nagradu za dobro obavljenu uslugu "
    "(dalje u tekstu: napojnica) od strane trećih osoba."
)
p = doc.add_paragraph(tekst)
p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.space_after = Pt(24)

# Signature/date block
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
left = p.add_run("Direktor:")
p.add_run("\t\t\t\t\t\t\t\t\t")
right = p.add_run("Datum:")

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.LEFT
p.add_run("____________________")
p.add_run("\t\t\t\t\t\t")
p.add_run("__________________")


doc.save(path)
print("OK: Styled DOCX saved", path)
